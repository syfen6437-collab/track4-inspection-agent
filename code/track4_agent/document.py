from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _read_optional(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8-sig") as handle:
        return json.load(handle)


def _set_run_font(run, name: str, size: float, *, bold: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _style_paragraph(paragraph, *, size: float = 10.5, bold: bool = False) -> None:
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.3
    for run in paragraph.runs:
        _set_run_font(run, "Arial", size, bold=bold)


def _insert_paragraph_after(paragraph, text: str = "", style: str | None = None):
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    new_element.getparent().replace(new_element, new_paragraph._p)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def _set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def _format_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_border(cell)
            if row_index == 0:
                _set_cell_fill(cell, "1D2330")
            elif row_index % 2 == 0:
                _set_cell_fill(cell, "F4F7FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    _set_run_font(
                        run,
                        "Arial",
                        8.5,
                        bold=row_index == 0,
                        color="FFFFFF" if row_index == 0 else "000000",
                    )


def _insert_table_after(paragraph, headers: list[str], rows: list[list[str]]):
    document = paragraph._parent
    table = document.add_table(rows=1, cols=len(headers), width=Inches(5.75))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    paragraph._p.addnext(table._tbl)
    _format_table(table)
    return table


def _find_paragraph(document: Document, startswith: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(startswith):
            return paragraph
    raise KeyError(startswith)


def build_design_document(workspace: Path, team_name: str = "") -> Path:
    template = workspace / "data" / "raw" / "AI智能体设计方案.docx"
    output = workspace / "design" / "AI智能体设计方案.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    if not template.exists():
        raise FileNotFoundError(template)

    dataset = _read_optional(workspace / "data" / "processed" / "dataset_summary.json")
    calibration = _read_optional(workspace / "logs" / "calibration_metrics.json")
    inference = _read_optional(workspace / "logs" / "inference_summary.json")
    validation = _read_optional(workspace / "logs" / "validation_report.json")

    document = Document(template)

    title = _find_paragraph(document, "AI智能体设计方案")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    metadata = {
        "参赛队伍": f"参赛队伍：{team_name}" if team_name else "参赛队伍：_________",
        "作品名称": "作品名称：城市路桥隧边坡结构病害智能巡检与分级评定智能体",
        "完成日期": f"完成日期：{date(2026, 10, 5).strftime('%Y年%m月%d日')}",
    }
    for prefix, value in metadata.items():
        paragraph = _find_paragraph(document, prefix)
        paragraph.text = value
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _style_paragraph(paragraph, size=10.5, bold=True)

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith(("一、", "二、", "三、", "四、", "五、")):
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_with_next = True
            if text.startswith("五、"):
                paragraph.paragraph_format.page_break_before = True
            for run in paragraph.runs:
                _set_run_font(run, "Arial", 16, bold=True)
        elif text.startswith(("1.", "2.")):
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                _set_run_font(run, "Arial", 10.5, bold=False)

    anchor = _find_paragraph(document, "一、需求与痛点分析")
    paragraphs = [
        "城市桥梁和轨道结构巡检图片数量大、拍摄视角差异明显，人工逐张检查耗时长，且病害类型、描述粒度和评定标度容易受到人员经验影响。本赛题数据同时包含航拍道路与桥梁全景、梁底和支座近景、轨道结构病害近景，存在明显的跨场景域差异。",
        "训练集共包含3,402张图片。桥梁样本以完好、渗水泛碱、修补痕迹和混凝土外观问题为主，轨道样本以裂缝、破损、支座锈蚀、钢筋锈蚀和钢结构锈蚀为主。系统需要在统一输出规范下完成病害识别、证据描述和分级评定，同时保证结果完全由模型和智能体自动生成。",
        "方案采用开源本地多模态模型，避免依赖外部接口，并通过训练标签词典约束、分域图像处理、局部裁剪复核和结构化结果校验，提高输出稳定性和工程可复现性。",
    ]
    for text in paragraphs:
        anchor = _insert_paragraph_after(anchor, text)
        _style_paragraph(anchor)

    anchor = _find_paragraph(document, "二、智能体总体架构设计")
    p = _insert_paragraph_after(
        anchor,
        "智能体由数据准备、场景路由、视觉推理、低置信度复核、结构化校验和提交打包六个模块组成。基座模型为Qwen3-VL-2B-Instruct，采用BF16在本地GPU运行；推理阶段不访问网络。",
    )
    _style_paragraph(p)
    architecture_rows = [
        ["数据准备", "公开压缩包、训练JSON", "安全解压、完全重复去重、图片与标注映射"],
        ["场景路由", "文件夹与照片编号", "区分桥梁、轨道、左右幅和连续航拍序列"],
        ["视觉推理", "桥梁整图或轨道整图加细节拼图", "输出病害类型、描述、标度、置信度和证据"],
        ["复核模块", "轨道低置信度结果", "使用整图和四象限细节拼图再次推理"],
        ["结果校验", "模型原始JSON", "合法标签规范化、字段检查、断点保存"],
        ["提交打包", "代码、设计书、结果", "生成符合赛事目录结构的tar.gz"],
    ]
    _insert_table_after(p, ["模块", "输入", "主要功能"], architecture_rows)

    anchor = _find_paragraph(document, "三、核心模块详细设计")
    module_texts = [
        "数据预处理模块：安全解压20GB公开数据，读取汇总标注，删除一条完全重复记录，并通过桥梁名称、左右幅和文件夹解决重复照片编号映射。生成训练清单、测试清单和标签词典，所有中间文件均可复核。",
        "分域视觉模块：六小时首版将桥梁整图最长边缩放到384像素；轨道整图保持448像素，并把四个重叠象限缩略图拼成一张2×2细节图。模型仍观察四个局部区域，但只编码两张图，兼顾细裂缝识别与吞吐。",
        "结构化推理模块：提示词只允许模型从训练集中出现的合法病害类型中选择。模型同时生成可见证据、病害描述和评定标度，温度设为0并固定随机种子，降低输出波动。",
        "自动复核模块：当轨道结果置信度低于0.65或完好标签与病害描述冲突时触发第二轮推理。六小时首版关闭桥梁二次复核，非法标签由训练集词典自动规范化，避免重复推理拖延完整提交。",
        "工程可靠性模块：每25张图片持久化一次断点，推理中断后可继续运行；原始模型响应保存在日志中。最终校验测试图片数量、文件名多重集合、字段顺序、合法标签和标度范围。",
    ]
    for text in module_texts:
        anchor = _insert_paragraph_after(anchor, text)
        _style_paragraph(anchor)

    anchor = _find_paragraph(document, "四、创新点设计")
    innovations = [
        "基于数据域的自适应视觉输入：桥梁采用轻量整图，轨道采用整图加四象限拼图，同一模型按场景切换视觉预算。",
        "训练集驱动的标签约束：标签、典型描述和评定标度分布全部从公开训练数据自动提取，既抑制大模型幻觉，又不包含任何测试答案。",
        "证据一致性复核：模型不仅给出类别，还给出可见依据；系统检测类别与描述冲突并自动复核，提高结构化结果可信度。",
        "全流程可追溯：保存模型版本、固定种子、每张图片原始响应、断点和校验报告，便于赛事复核和决赛演示。",
    ]
    for index, text in enumerate(innovations, start=1):
        anchor = _insert_paragraph_after(anchor, f"{index}. {text}")
        _style_paragraph(anchor)

    evaluation_heading = _find_paragraph(document, "1. 评估结果")
    metrics_rows = [
        ["训练图片", str(dataset.get("train_images", "待生成")), "公开数据完整性检查"],
        ["测试图片", str(dataset.get("test_images", "待生成")), "最终结果应逐图覆盖"],
        ["校准样本", str(calibration.get("sample_size", 0)), "覆盖低频标签的压力抽样"],
        ["病害类型精确率", f"{calibration.get('defect_type_exact_accuracy', 0):.4f}", "仅用于本地版本比较"],
        ["原子病害Macro-F1", f"{calibration.get('atomic_macro_f1', 0):.4f}", "组合病害拆分后计算"],
        ["评定标度准确率", f"{calibration.get('rating_accuracy', 0):.4f}", "含空标度"],
        ["结果文件校验", "通过" if validation.get("valid") else "待校验", "七字段、数量及标签合法性"],
    ]
    _insert_table_after(evaluation_heading, ["指标", "结果", "说明"], metrics_rows)

    time_heading = _find_paragraph(document, "2. 任务执行耗时")
    elapsed = float(inference.get("elapsed_seconds", 0) or 0)
    time_rows = [
        ["数据准备", "自动记录", "解压、去重、清单与词典生成"],
        ["模型推理", f"{elapsed / 3600:.2f}小时", f"共完成{inference.get('completed', 0)}张测试图片"],
        ["运行吞吐", f"{inference.get('images_per_second', 0)}张/秒", "包含模型生成及必要复核"],
        ["结果校验与打包", "自动执行", "生成校验报告和tar.gz提交包"],
    ]
    _insert_table_after(time_heading, ["阶段", "耗时或结果", "内容"], time_rows)

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            if not paragraph.style.name.startswith("Title"):
                for run in paragraph.runs:
                    if run.font.size is None:
                        _set_run_font(run, "Arial", 10.5)

    core = document.core_properties
    core.title = "城市路桥隧边坡结构病害智能巡检与分级评定智能体设计方案"
    core.subject = "第二届国际通用人工智能大会行业智能体创新挑战赛赛道四"
    core.author = team_name
    document.save(output)
    print(f"[document] saved {output}")
    return output
