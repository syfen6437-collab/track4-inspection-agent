from __future__ import annotations

import hashlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from track4_agent.document import refresh_submission_design


class DocumentRefreshTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        for name in ("design", "logs", "result"):
            (self.root / name).mkdir()
        self.path = self.root / "design" / "AI智能体设计方案.docx"
        document = Document()
        for text in ("AI智能体设计方案", "参赛队伍：保留队伍", "完成日期：2026年10月05日",
                     "工程可靠性模块：旧说明", "1. 基于数据域：旧说明", "3. 证据一致性：旧说明",
                     "分域视觉模块：旧说明", "1. 评估结果"):
            document.add_paragraph(text)
        table = document.add_table(rows=1, cols=3)
        for cell, text in zip(table.rows[0].cells, ("指标", "值", "说明")):
            cell.text = text
        for values in (("校准样本", "24", "旧说明"), ("模型推理", "1.47小时", "旧说明")):
            for cell, text in zip(table.add_row().cells, values):
                cell.text = text
        document.save(self.path)
        self.original = self.path.read_bytes()
        self.result = self.root / "result" / "result.json"
        self.result.write_bytes(b"[]")
        self.provenance = {"status": "restored_scored_baseline", "source_commit": "d2f6bde-example",
                           "result_sha256": hashlib.sha256(b"[]").hexdigest(),
                           "official_feedback": {"total": 56.11}}
        self.provenance_path = self.root / "logs" / "submission_provenance.json"
        self.provenance_path.write_text(json.dumps(self.provenance), encoding="utf-8")
        (self.root / "logs" / "validation_report.json").write_text(
            json.dumps({"valid": True, "result_count": 0}), encoding="utf-8")

    def test_refresh_preserves_team_result_and_all_unrelated_package_parts(self):
        self.assertEqual(refresh_submission_design(self.root), self.path)
        document = Document(self.path)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("参赛队伍：保留队伍", text)
        self.assertIn("56.11", text)
        self.assertIn("未完成两座整桥", text)
        self.assertEqual(self.result.read_bytes(), b"[]")
        with zipfile.ZipFile(io.BytesIO(self.original)) as old, zipfile.ZipFile(self.path) as new:
            self.assertEqual(old.namelist(), new.namelist())
            for name in old.namelist():
                if name != "word/document.xml":
                    self.assertEqual(old.read(name), new.read(name), name)

    def test_repeated_refresh_does_not_duplicate_note_or_table_header(self):
        refresh_submission_design(self.root)
        refresh_submission_design(self.root)
        document = Document(self.path)
        self.assertEqual(sum(p.text.startswith("提交版本说明：") for p in document.paragraphs), 1)
        from docx.oxml.ns import qn
        self.assertEqual(len(document.tables[0].rows[0]._tr.trPr.findall(qn("w:tblHeader"))), 1)

    def test_changed_result_is_rejected_without_editing_document(self):
        self.result.write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "Validate and register"):
            refresh_submission_design(self.root)
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_unregistered_candidate_cannot_refresh_baseline_design(self):
        self.provenance["status"] = "candidate"
        self.provenance_path.write_text(json.dumps(self.provenance), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "historical baseline"):
            refresh_submission_design(self.root)
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_completed_single_fold_is_used_instead_of_early_pressure_sample(self):
        folder = self.root / "runs" / "prompt_ablation_v3"
        folder.mkdir(parents=True)
        metrics = {"sample_size": 120, "defect_type_exact_accuracy": 38 / 120,
                   "atomic_macro_f1": 0.25, "rating_accuracy": 0.55,
                   "category_metrics": {"桥梁": {"exact_correct": 33}, "轨道": {"exact_correct": 5}}}
        (folder / "fanjaping.json").write_text(json.dumps({"results": {"baseline": {
            "status": "complete", "records": [{} for _ in range(120)], "metrics": metrics}}}), encoding="utf-8")
        refresh_submission_design(self.root)
        document = Document(self.path)
        row = document.tables[0].rows[1]
        self.assertEqual(row.cells[0].text, "留出样本")
        self.assertEqual(row.cells[1].text, "120")
        self.assertIn("单折基线", row.cells[-1].text)


if __name__ == "__main__":
    unittest.main()
